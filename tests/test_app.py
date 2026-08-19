from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import create_app


@pytest.fixture()
def app(tmp_path: Path):
    db = tmp_path / "test.db"
    application = create_app(
        {
            "TESTING": True,
            "SECRET_KEY": "test-secret",
            "DATABASE": str(db),
            "DATA_DIR": str(tmp_path),
            "PHOTOS_DIR": str(tmp_path / "photos"),
            "ADMIN_EMAIL": "admin@saintmarkpathum.com",
            "ADMIN_PASSWORD": "admin123",
        }
    )
    (tmp_path / "photos").mkdir(exist_ok=True)
    return application


@pytest.fixture()
def client(app):
    return app.test_client()


def login(client, email="admin@saintmarkpathum.com", password="admin123"):
    return client.post("/api/login", json={"email": email, "password": password})


def test_health(client):
    res = client.get("/health")
    assert res.status_code == 200
    assert res.get_json()["ok"] is True


def test_login_pages(client):
    assert client.get("/").status_code == 200
    assert "ระบบรับของบริจาคจาก 7-11" in client.get("/").get_data(as_text=True)
    assert client.get("/donate").status_code == 200
    assert client.get("/reset-password").status_code == 200


def test_login_required(client):
    assert client.get("/api/branches").status_code == 401
    assert login(client, password="wrong").status_code == 401
    assert login(client).status_code == 200
    assert client.get("/api/me").get_json()["user"]["email"] == "admin@saintmarkpathum.com"


def test_seeded_branches(client):
    login(client)
    data = client.get("/api/branches").get_json()
    codes = {b["code"] for b in data["branches"]}
    assert "15005" in codes
    assert "11583" in codes
    assert len(data["branches"]) >= 38


def test_public_branch_lookup(client):
    res = client.get("/api/public/branch?code=15005")
    assert res.status_code == 200
    assert res.get_json()["branch"]["name"] == "หมู่บ้านเมืองเอก-รังสิต"
    # leading zeros / stripped zeros
    res = client.get("/api/public/branch?code=2196")
    assert res.status_code == 200
    assert res.get_json()["branch"]["code"] == "02196"
    assert client.get("/api/public/branch?code=nope").status_code == 404


def test_public_and_staff_donation_flow(client):
    res = client.post(
        "/api/public/donations",
        json={
            "date": "2026-08-19",
            "branchCode": "15005",
            "pieces": 12,
            "weightKg": 3.5,
            "baskets": 2,
            "contactName": "สมชาย",
            "phone": "0812345678",
        },
    )
    assert res.status_code == 200
    don = res.get_json()["donation"]
    assert don["storeName"] == "หมู่บ้านเมืองเอก-รังสิต"
    assert don["source"] == "public-form"

    login(client)
    staff = client.post(
        "/api/donations",
        json={"date": "2026-08-19", "branchCode": "11583", "pieces": 5, "weightKg": 1},
    )
    assert staff.status_code == 200

    listed = client.get("/api/donations?mode=day&date=2026-08-19").get_json()
    assert listed["totals"]["totalPieces"] == 17
    assert listed["totals"]["totalRecords"] == 2

    summary = client.get("/api/donations/month-summary?year=2026&month=8").get_json()
    assert "สิงหาคม" in summary["label"]
    assert summary["totals"]["totalPieces"] == 17

    deleted = client.delete(f"/api/donations/{don['id']}")
    assert deleted.status_code == 200
    listed = client.get("/api/donations?mode=day&date=2026-08-19").get_json()
    assert listed["totals"]["totalRecords"] == 1


def test_groups_and_reports(client):
    login(client)
    client.post("/api/donations", json={"date": "2026-08-01", "branchCode": "15005", "pieces": 10})
    client.post("/api/donations", json={"date": "2026-08-01", "branchCode": "2664", "pieces": 4})
    created = client.post(
        "/api/groups",
        json={"name": "สายเมืองเอก", "description": "รับของทุกวัน", "members": ["15005", "2664"]},
    )
    assert created.status_code == 200
    groups = client.get("/api/groups").get_json()["groups"]
    assert groups[0]["name"] == "สายเมืองเอก"
    assert set(groups[0]["members"]) == {"15005", "2664"}

    report = client.get("/api/reports/by-group?mode=day&date=2026-08-01").get_json()
    found = [g for g in report["groups"] if g["groupName"] == "สายเมืองเอก"][0]
    assert found["totalPieces"] == 14

    branch = client.get("/api/reports/by-branch?year=2026&month=8").get_json()
    assert branch["totals"]["branchCount"] == 2


def test_line_groups_unique(client):
    login(client)
    ok = client.post(
        "/api/line-groups",
        json={"name": "กลุ่มหลัก", "group_id": "Cabc", "message_type": "full"},
    )
    assert ok.status_code == 200
    dup = client.post(
        "/api/line-groups",
        json={"name": "ซ้ำ", "group_id": "Cabc", "message_type": "short"},
    )
    assert dup.status_code == 400
    webhook = client.post(
        "/api/line/webhook",
        json={"events": [{"source": {"groupId": "Cnewgroup"}}]},
    )
    assert webhook.status_code == 200
    groups = client.get("/api/line-groups").get_json()["groups"]
    ids = {g["group_id"] for g in groups}
    assert "Cabc" in ids and "Cnewgroup" in ids


def test_add_branch_and_contact(client):
    login(client)
    res = client.post("/api/branches", json={"code": "99999", "name": "สาขาทดสอบ", "contactName": "นิด"})
    assert res.status_code == 200
    again = client.post("/api/branches", json={"code": "99999", "name": "ซ้ำ"})
    assert again.status_code == 400
    client.put("/api/branch-contacts/99999", json={"contactName": "ใหม่", "phone": "02"})
    data = client.get("/api/branches").get_json()
    row = next(b for b in data["branches"] if b["code"] == "99999")
    assert row["contactName"] == "ใหม่"


def test_excel_export(client):
    login(client)
    client.post("/api/donations", json={"date": "2026-08-19", "branchCode": "15005", "pieces": 8, "weightKg": 2})
    res = client.get("/api/export/donations.xlsx?mode=day&date=2026-08-19")
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.data))
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0][0] == "วันที่"
    assert any(r[1] == "15005" for r in rows[1:])


def test_forgot_and_reset_password(client):
    res = client.post("/api/forgot-password", json={"email": "admin@saintmarkpathum.com"})
    assert res.status_code == 200
    token = res.get_json()["token"]
    bad = client.post("/api/reset-password", json={"token": token, "password": "123", "confirm": "123"})
    assert bad.status_code == 400
    ok = client.post(
        "/api/reset-password",
        json={"token": token, "password": "newpass", "confirm": "newpass"},
    )
    assert ok.status_code == 200
    assert login(client, password="admin123").status_code == 401
    assert login(client, password="newpass").status_code == 200


def test_photo_upload(client, tmp_path):
    login(client)
    data = {
        "date": "2026-08-19",
        "branchCode": "15005",
        "pieces": "3",
        "photos": (io.BytesIO(b"\x89PNG\r\n\x1a\n" + b"0" * 20), "shot.png"),
    }
    res = client.post("/api/donations", data=data, content_type="multipart/form-data")
    assert res.status_code == 200
    photos = res.get_json()["photos"]
    assert photos
    got = client.get(f"/photos/{photos[0]}")
    assert got.status_code == 200


def test_import_saintmark_snapshot(tmp_path: Path):
    import db as dbmod

    conn = dbmod.connect(tmp_path / "import.db")
    dbmod.init_db(conn, "admin@saintmarkpathum.com", "admin123")
    stats = dbmod.import_snapshot(conn)
    assert stats["donations"] == 3129
    assert conn.execute("SELECT COUNT(*) FROM donations").fetchone()[0] == 3129
    assert conn.execute("SELECT COUNT(*) FROM line_groups").fetchone()[0] == 5
    assert conn.execute("SELECT COUNT(*) FROM branches").fetchone()[0] >= 43
    row = conn.execute(
        "SELECT pieces, store_name FROM donations WHERE branch_code=? ORDER BY created_at LIMIT 1",
        ("11433",),
    ).fetchone()
    assert row["store_name"] == "STUDENT CENTER ม.รังสิต"
    # second import must not duplicate
    dbmod.import_snapshot(conn)
    assert conn.execute("SELECT COUNT(*) FROM donations").fetchone()[0] == 3129
    pieces = conn.execute("SELECT SUM(pieces) FROM donations").fetchone()[0]
    assert pieces == 53566
