from __future__ import annotations

import io
import os
import secrets
import sqlite3
import uuid
from datetime import datetime, timedelta, timezone
from functools import wraps
from pathlib import Path
from typing import Any, Callable
from urllib.parse import urljoin

from flask import (
    Flask,
    Response,
    g,
    jsonify,
    redirect,
    render_template,
    request,
    send_from_directory,
    session,
)
from openpyxl import Workbook
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

import db as dbmod
from line_notify import notify_line_groups

THAI_MONTHS = [
    "มกราคม",
    "กุมภาพันธ์",
    "มีนาคม",
    "เมษายน",
    "พฤษภาคม",
    "มิถุนายน",
    "กรกฎาคม",
    "สิงหาคม",
    "กันยายน",
    "ตุลาคม",
    "พฤศจิกายน",
    "ธันวาคม",
]

def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def iso_today() -> str:
    return datetime.now().date().isoformat()


def month_label(year: int, month: int) -> str:
    return f"{THAI_MONTHS[month - 1]} {year + 543}"


def create_app(test_config: dict[str, Any] | None = None) -> Flask:
    app = Flask(__name__, static_folder="static", template_folder="templates")
    root = Path(__file__).resolve().parent
    data_dir = Path(os.environ.get("GIFT_DATA_DIR", root / "data"))
    data_dir.mkdir(parents=True, exist_ok=True)
    photos_dir = data_dir / "photos"
    photos_dir.mkdir(parents=True, exist_ok=True)

    secret_file = data_dir / ".secret"
    if test_config and "SECRET_KEY" in test_config:
        secret = test_config["SECRET_KEY"]
    elif os.environ.get("GIFT_SECRET_KEY"):
        secret = os.environ["GIFT_SECRET_KEY"]
    elif secret_file.exists():
        secret = secret_file.read_text().strip()
    else:
        secret = secrets.token_hex(32)
        secret_file.write_text(secret)

    db_path = Path(os.environ.get("GIFT_DB", data_dir / "gift_logger.db"))
    if test_config and "DATABASE" in test_config:
        db_path = Path(test_config["DATABASE"])

    app.config.update(
        SECRET_KEY=secret,
        DATABASE=str(db_path),
        DATA_DIR=str(data_dir),
        PHOTOS_DIR=str(photos_dir),
        ADMIN_EMAIL=os.environ.get("GIFT_ADMIN_EMAIL", "admin@saintmarkpathum.com"),
        ADMIN_PASSWORD=os.environ.get("GIFT_ADMIN_PASSWORD", "admin123"),
        LINE_CHANNEL_ACCESS_TOKEN=os.environ.get("LINE_CHANNEL_ACCESS_TOKEN", ""),
        LINE_CHANNEL_SECRET=os.environ.get("LINE_CHANNEL_SECRET", ""),
        PUBLIC_BASE_URL=os.environ.get("PUBLIC_BASE_URL", "").rstrip("/"),
        MAX_CONTENT_LENGTH=16 * 1024 * 1024,
        SESSION_COOKIE_SAMESITE="Lax",
    )
    if test_config:
        app.config.update(test_config)
    Path(app.config["PHOTOS_DIR"]).mkdir(parents=True, exist_ok=True)
    Path(app.config["DATA_DIR"]).mkdir(parents=True, exist_ok=True)

    def get_db() -> sqlite3.Connection:
        if "db" not in g:
            g.db = dbmod.connect(Path(app.config["DATABASE"]))
        return g.db

    @app.teardown_appcontext
    def close_db(_exc: BaseException | None) -> None:
        conn = g.pop("db", None)
        if conn is not None:
            conn.close()

    with app.app_context():
        conn = get_db()
        dbmod.init_db(conn, app.config["ADMIN_EMAIL"], app.config["ADMIN_PASSWORD"])
        skip_import = bool(app.config.get("TESTING") or os.environ.get("GIFT_SKIP_IMPORT"))
        if not skip_import and conn.execute("SELECT COUNT(*) FROM donations").fetchone()[0] == 0:
            dbmod.import_snapshot(conn)

    def login_required(fn: Callable) -> Callable:
        @wraps(fn)
        def wrapper(*args: Any, **kwargs: Any):
            if not session.get("user_id"):
                return jsonify({"error": "unauthorized"}), 401
            return fn(*args, **kwargs)

        return wrapper

    def current_user() -> sqlite3.Row | None:
        uid = session.get("user_id")
        if not uid:
            return None
        return get_db().execute("SELECT id, email FROM users WHERE id=?", (uid,)).fetchone()

    def row_donation(r: sqlite3.Row) -> dict[str, Any]:
        return {
            "id": r["id"],
            "date": r["date"],
            "pickupDate": r["pickup_date"],
            "storeName": r["store_name"],
            "branchCode": r["branch_code"],
            "pieces": r["pieces"],
            "weightKg": float(r["weight_kg"] or 0),
            "baskets": r["baskets"],
            "contactName": r["contact_name"],
            "position": r["position"],
            "phone": r["phone"],
            "source": r["source"],
            "createdAt": r["created_at"],
        }

    def list_branches() -> list[dict[str, str]]:
        rows = get_db().execute(
            "SELECT code, name FROM branches ORDER BY name COLLATE NOCASE"
        ).fetchall()
        return [{"code": r["code"], "name": r["name"]} for r in rows]

    def find_branch(code: str) -> dict[str, str] | None:
        raw = (code or "").strip()
        if not raw:
            return None
        row = get_db().execute("SELECT code, name FROM branches WHERE code=?", (raw,)).fetchone()
        if row:
            return {"code": row["code"], "name": row["name"]}
        stripped = raw.lstrip("0") or "0"
        rows = get_db().execute("SELECT code, name FROM branches").fetchall()
        for r in rows:
            if (r["code"].lstrip("0") or "0") == stripped:
                return {"code": r["code"], "name": r["name"]}
        return None

    def donations_in_range(start: str, end: str) -> list[dict[str, Any]]:
        rows = get_db().execute(
            """
            SELECT * FROM donations
            WHERE pickup_date >= ? AND pickup_date <= ?
            ORDER BY branch_code, pickup_date, id
            """,
            (start, end),
        ).fetchall()
        return [row_donation(r) for r in rows]

    def month_bounds(year: int, month: int) -> tuple[str, str]:
        start = f"{year:04d}-{month:02d}-01"
        if month == 12:
            end_excl = f"{year + 1:04d}-01-01"
        else:
            end_excl = f"{year:04d}-{month + 1:02d}-01"
        return start, end_excl

    def donations_in_month(year: int, month: int) -> list[dict[str, Any]]:
        start, end_excl = month_bounds(year, month)
        rows = get_db().execute(
            """
            SELECT * FROM donations
            WHERE pickup_date >= ? AND pickup_date < ?
            ORDER BY branch_code, pickup_date, id
            """,
            (start, end_excl),
        ).fetchall()
        return [row_donation(r) for r in rows]

    def load_groups() -> list[dict[str, Any]]:
        conn = get_db()
        groups = conn.execute(
            "SELECT * FROM branch_groups ORDER BY created_at, id"
        ).fetchall()
        members = conn.execute("SELECT group_id, branch_code FROM branch_group_members").fetchall()
        by_g: dict[int, list[str]] = {}
        for m in members:
            by_g.setdefault(m["group_id"], []).append(m["branch_code"])
        return [
            {
                "id": g["id"],
                "name": g["name"],
                "description": g["description"],
                "members": by_g.get(g["id"], []),
            }
            for g in groups
        ]

    def save_photos(donation_id: str, files: list[Any]) -> list[str]:
        saved: list[str] = []
        dest = Path(app.config["PHOTOS_DIR"]) / str(donation_id)
        dest.mkdir(parents=True, exist_ok=True)
        conn = get_db()
        for i, f in enumerate(files):
            if not f or not getattr(f, "filename", None):
                continue
            ext = Path(secure_filename(f.filename)).suffix.lower() or ".jpg"
            if ext not in {".jpg", ".jpeg", ".png", ".gif", ".webp"}:
                ext = ".jpg"
            name = f"{secrets.token_hex(8)}{ext}"
            path = dest / name
            f.save(path)
            rel = f"{donation_id}/{name}"
            conn.execute(
                "INSERT INTO donation_photos(donation_id, storage_path) VALUES (?,?)",
                (donation_id, rel),
            )
            saved.append(rel)
        conn.commit()
        return saved

    def insert_donation(payload: dict[str, Any], source: str) -> dict[str, Any] | None:
        branch = find_branch(str(payload.get("branchCode") or payload.get("branch_code") or ""))
        if not branch:
            return None
        date = (payload.get("date") or iso_today()).strip()
        try:
            pieces = int(payload.get("pieces") or 0)
        except (TypeError, ValueError):
            pieces = 0
        if pieces <= 0:
            return None
        try:
            weight = float(payload.get("weightKg") or payload.get("weight_kg") or 0)
        except (TypeError, ValueError):
            weight = 0
        baskets = payload.get("baskets")
        try:
            baskets_i = int(baskets) if baskets not in (None, "") else None
        except (TypeError, ValueError):
            baskets_i = None
        now = utc_now()
        donation_id = str(uuid.uuid4())
        conn = get_db()
        conn.execute(
            """
            INSERT INTO donations(
                id, date, pickup_date, store_name, branch_code, pieces, weight_kg,
                baskets, contact_name, position, phone, source, created_at
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                donation_id,
                date,
                date,
                branch["name"],
                branch["code"],
                pieces,
                weight,
                baskets_i,
                (payload.get("contactName") or payload.get("contact_name") or None) or None,
                (payload.get("position") or None) or None,
                (payload.get("phone") or None) or None,
                source,
                now,
            ),
        )
        conn.commit()
        row = conn.execute("SELECT * FROM donations WHERE id=?", (donation_id,)).fetchone()
        return row_donation(row)

    def public_photo_urls(paths: list[str]) -> list[str]:
        base = app.config["PUBLIC_BASE_URL"]
        if not base:
            return []
        return [urljoin(base + "/", f"photos/{p}") for p in paths]

    def notify(donation: dict[str, Any], photo_paths: list[str]) -> None:
        token = app.config.get("LINE_CHANNEL_ACCESS_TOKEN") or ""
        if not token:
            return
        groups = [
            dict(r)
            for r in get_db().execute("SELECT * FROM line_groups ORDER BY id").fetchall()
        ]
        notify_line_groups(
            token,
            groups,
            {
                "date": donation["date"],
                "branchCode": donation["branchCode"],
                "storeName": donation["storeName"],
                "pieces": donation["pieces"],
                "baskets": donation["baskets"],
                "contactName": donation["contactName"],
                "position": donation["position"],
                "phone": donation["phone"],
            },
            public_photo_urls(photo_paths),
        )

    def xlsx_response(wb: Workbook, filename: str) -> Response:
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return Response(
            buf.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    # ---------- pages ----------
    @app.get("/")
    def index():
        return render_template("app.html", page="home")

    @app.get("/donate")
    def donate_page():
        return render_template("app.html", page="donate")

    @app.get("/reset-password")
    def reset_page():
        return render_template("app.html", page="reset")

    @app.get("/health")
    def health():
        return jsonify({"ok": True, "service": "seven-gift-logger"})

    @app.get("/photos/<path:rel>")
    def photos(rel: str):
        return send_from_directory(app.config["PHOTOS_DIR"], rel)

    # ---------- auth ----------
    @app.get("/api/me")
    def me():
        user = current_user()
        if not user:
            return jsonify({"user": None})
        return jsonify({"user": {"id": user["id"], "email": user["email"]}})

    @app.post("/api/login")
    def login():
        body = request.get_json(silent=True) or {}
        email = (body.get("email") or "").strip().lower()
        password = body.get("password") or ""
        if not email:
            return jsonify({"error": "กรุณากรอกอีเมล"}), 400
        if not password:
            return jsonify({"error": "กรุณากรอกรหัสผ่าน"}), 400
        row = get_db().execute("SELECT * FROM users WHERE email=?", (email,)).fetchone()
        if not row or not check_password_hash(row["password_hash"], password):
            return jsonify({"error": "เกิดข้อผิดพลาด"}), 401
        session["user_id"] = row["id"]
        session.permanent = True
        return jsonify({"ok": True, "user": {"id": row["id"], "email": row["email"]}})

    @app.post("/api/logout")
    def logout():
        session.clear()
        return jsonify({"ok": True})

    @app.post("/api/forgot-password")
    def forgot_password():
        body = request.get_json(silent=True) or {}
        email = (body.get("email") or "").strip().lower()
        if not email:
            return jsonify({"error": "กรุณากรอกอีเมล"}), 400
        row = get_db().execute("SELECT id FROM users WHERE email=?", (email,)).fetchone()
        # Always look successful to match typical auth UX, but store a token when user exists.
        if row:
            token = secrets.token_urlsafe(24)
            expires = (datetime.now(timezone.utc) + timedelta(hours=2)).isoformat()
            get_db().execute(
                "INSERT OR REPLACE INTO password_resets(token, email, expires_at) VALUES (?,?,?)",
                (token, email, expires),
            )
            get_db().commit()
            app.logger.info("Password reset token for %s: %s", email, token)
            if app.config.get("TESTING"):
                return jsonify({"ok": True, "token": token})
        return jsonify({"ok": True, "message": "ส่งลิงก์รีเซ็ตรหัสผ่านไปยังอีเมลแล้ว"})

    @app.post("/api/reset-password")
    def reset_password():
        body = request.get_json(silent=True) or {}
        token = (body.get("token") or "").strip()
        password = body.get("password") or ""
        confirm = body.get("confirm") or password
        if len(password) < 6:
            return jsonify({"error": "รหัสผ่านต้องมีอย่างน้อย 6 ตัวอักษร"}), 400
        if password != confirm:
            return jsonify({"error": "รหัสผ่านทั้งสองช่องไม่ตรงกัน"}), 400
        row = get_db().execute(
            "SELECT * FROM password_resets WHERE token=?", (token,)
        ).fetchone()
        if not row or row["expires_at"] < utc_now():
            return jsonify({"error": "ลิงก์รีเซ็ตไม่ถูกต้องหรือหมดอายุ"}), 400
        get_db().execute(
            "UPDATE users SET password_hash=? WHERE email=?",
            (generate_password_hash(password), row["email"]),
        )
        get_db().execute("DELETE FROM password_resets WHERE token=?", (token,))
        get_db().commit()
        return jsonify({"ok": True})

    # ---------- public donate ----------
    @app.get("/api/public/branch")
    def public_branch():
        code = request.args.get("code") or ""
        branch = find_branch(code)
        if not branch:
            return jsonify({"error": "ไม่พบรหัสสาขานี้ กรุณาตรวจสอบอีกครั้ง"}), 404
        contacts = get_db().execute(
            "SELECT contact_name, phone, position FROM branch_contacts WHERE branch_code=?",
            (branch["code"],),
        ).fetchall()
        return jsonify(
            {
                "branch": branch,
                "contacts": [dict(c) for c in contacts if c["contact_name"] or c["phone"]],
            }
        )

    @app.get("/api/public/branches")
    def public_branches():
        return jsonify({"branches": list_branches()})

    @app.post("/api/public/donations")
    def public_create_donation():
        if request.is_json:
            payload = request.get_json() or {}
            files = []
        else:
            payload = {k: request.form.get(k) for k in request.form}
            files = request.files.getlist("photos")
        rec = insert_donation(payload, source="public-form")
        if not rec:
            return jsonify({"error": "บันทึกไม่สำเร็จ กรุณาลองใหม่"}), 400
        photos = save_photos(rec["id"], files)
        notify(rec, photos)
        return jsonify({"ok": True, "donation": rec, "photos": photos})

    # ---------- branches / contacts ----------
    @app.get("/api/branches")
    @login_required
    def api_branches():
        contacts = {
            r["branch_code"]: dict(r)
            for r in get_db().execute("SELECT * FROM branch_contacts").fetchall()
        }
        items = []
        for b in list_branches():
            c = contacts.get(b["code"], {})
            items.append(
                {
                    **b,
                    "contactName": c.get("contact_name"),
                    "phone": c.get("phone"),
                    "position": c.get("position"),
                }
            )
        return jsonify({"branches": items})

    @app.post("/api/branches")
    @login_required
    def api_add_branch():
        body = request.get_json(silent=True) or {}
        code = (body.get("code") or "").strip()
        name = (body.get("name") or "").strip()
        if not code or not name:
            return jsonify({"error": "กรุณากรอกรหัสร้านและชื่อสาขา"}), 400
        conn = get_db()
        if conn.execute("SELECT 1 FROM branches WHERE code=?", (code,)).fetchone():
            return jsonify({"error": "รหัสร้านนี้มีอยู่แล้ว"}), 400
        now = utc_now()
        conn.execute(
            "INSERT INTO branches(code, name, created_at) VALUES (?,?,?)",
            (code, name, now),
        )
        conn.execute(
            """
            INSERT INTO branch_contacts(branch_code, branch_name, contact_name, phone, position, updated_at)
            VALUES (?,?,?,?,?,?)
            """,
            (
                code,
                name,
                (body.get("contactName") or None),
                (body.get("phone") or None),
                (body.get("position") or None),
                now,
            ),
        )
        conn.commit()
        return jsonify({"ok": True, "branch": {"code": code, "name": name}})

    @app.put("/api/branch-contacts/<code>")
    @login_required
    def api_save_contact(code: str):
        body = request.get_json(silent=True) or {}
        now = utc_now()
        get_db().execute(
            """
            INSERT INTO branch_contacts(branch_code, branch_name, contact_name, phone, position, updated_at)
            VALUES (?,?,?,?,?,?)
            ON CONFLICT(branch_code) DO UPDATE SET
                contact_name=excluded.contact_name,
                phone=excluded.phone,
                position=excluded.position,
                updated_at=excluded.updated_at
            """,
            (
                code,
                body.get("branchName"),
                body.get("contactName") or None,
                body.get("phone") or None,
                body.get("position") or None,
                now,
            ),
        )
        get_db().commit()
        return jsonify({"ok": True})

    # ---------- donations ----------
    @app.get("/api/donations")
    @login_required
    def api_donations():
        mode = request.args.get("mode") or "day"
        if mode == "month":
            year = int(request.args.get("year") or datetime.now().year)
            month = int(request.args.get("month") or datetime.now().month)
            items = donations_in_month(year, month)
        elif mode == "range":
            start = request.args.get("start") or iso_today()
            end = request.args.get("end") or start
            items = donations_in_range(start, end)
        else:
            day = request.args.get("date") or iso_today()
            items = donations_in_range(day, day)
        return jsonify({"donations": items, "totals": totals_of(items)})

    @app.get("/api/donations/month-summary")
    @login_required
    def api_month_summary():
        year = int(request.args.get("year") or datetime.now().year)
        month = int(request.args.get("month") or datetime.now().month)
        items = donations_in_month(year, month)
        return jsonify(
            {
                "year": year,
                "month": month,
                "label": month_label(year, month),
                "totals": totals_of(items),
            }
        )

    @app.post("/api/donations")
    @login_required
    def api_create_donation():
        if request.is_json:
            payload = request.get_json() or {}
            files = []
        else:
            payload = {k: request.form.get(k) for k in request.form}
            files = request.files.getlist("photos")
        rec = insert_donation(payload, source=payload.get("source") or "web")
        if not rec:
            return jsonify({"error": "กรุณากรอกข้อมูลให้ครบ"}), 400
        photos = save_photos(rec["id"], files)
        notify(rec, photos)
        return jsonify({"ok": True, "donation": rec, "photos": photos})

    @app.delete("/api/donations/<donation_id>")
    @login_required
    def api_delete_donation(donation_id: str):
        conn = get_db()
        conn.execute("DELETE FROM donation_photos WHERE donation_id=?", (donation_id,))
        conn.execute("DELETE FROM donations WHERE id=?", (donation_id,))
        conn.commit()
        dest = Path(app.config["PHOTOS_DIR"]) / str(donation_id)
        if dest.exists():
            for p in dest.iterdir():
                p.unlink(missing_ok=True)
            dest.rmdir()
        return jsonify({"ok": True})

    # ---------- groups ----------
    @app.get("/api/groups")
    @login_required
    def api_groups():
        return jsonify({"groups": load_groups()})

    @app.post("/api/groups")
    @login_required
    def api_create_group():
        body = request.get_json(silent=True) or {}
        name = (body.get("name") or "").strip()
        if not name:
            return jsonify({"error": "กรุณาใส่ชื่อกลุ่ม"}), 400
        now = utc_now()
        conn = get_db()
        cur = conn.execute(
            "INSERT INTO branch_groups(name, description, created_at, updated_at) VALUES (?,?,?,?)",
            (name, (body.get("description") or None), now, now),
        )
        gid = cur.lastrowid
        for code in body.get("members") or []:
            conn.execute(
                "INSERT OR IGNORE INTO branch_group_members(group_id, branch_code) VALUES (?,?)",
                (gid, code),
            )
        conn.commit()
        return jsonify({"ok": True, "id": gid})

    @app.put("/api/groups/<int:gid>")
    @login_required
    def api_update_group(gid: int):
        body = request.get_json(silent=True) or {}
        name = (body.get("name") or "").strip()
        if not name:
            return jsonify({"error": "กรุณาใส่ชื่อกลุ่ม"}), 400
        conn = get_db()
        conn.execute(
            "UPDATE branch_groups SET name=?, description=?, updated_at=? WHERE id=?",
            (name, body.get("description") or None, utc_now(), gid),
        )
        conn.execute("DELETE FROM branch_group_members WHERE group_id=?", (gid,))
        for code in body.get("members") or []:
            conn.execute(
                "INSERT OR IGNORE INTO branch_group_members(group_id, branch_code) VALUES (?,?)",
                (gid, code),
            )
        conn.commit()
        return jsonify({"ok": True})

    @app.delete("/api/groups/<int:gid>")
    @login_required
    def api_delete_group(gid: int):
        conn = get_db()
        conn.execute("DELETE FROM branch_group_members WHERE group_id=?", (gid,))
        conn.execute("DELETE FROM branch_groups WHERE id=?", (gid,))
        conn.commit()
        return jsonify({"ok": True})

    # ---------- LINE groups ----------
    @app.get("/api/line-groups")
    @login_required
    def api_line_groups():
        rows = get_db().execute("SELECT * FROM line_groups ORDER BY created_at, id").fetchall()
        return jsonify({"groups": [dict(r) for r in rows]})

    @app.post("/api/line-groups")
    @login_required
    def api_add_line_group():
        body = request.get_json(silent=True) or {}
        name = (body.get("name") or "").strip()
        group_id = (body.get("group_id") or body.get("groupId") or "").strip()
        message_type = body.get("message_type") or body.get("messageType") or "full"
        if message_type not in ("full", "short", "summary"):
            message_type = "full"
        if not name or not group_id:
            return jsonify({"error": "กรุณากรอกข้อมูลให้ครบ"}), 400
        try:
            get_db().execute(
                "INSERT INTO line_groups(id, name, group_id, message_type, created_at) VALUES (?,?,?,?,?)",
                (str(uuid.uuid4()), name, group_id, message_type, utc_now()),
            )
            get_db().commit()
        except sqlite3.IntegrityError:
            return jsonify({"error": "Group ID นี้มีอยู่แล้ว"}), 400
        return jsonify({"ok": True})

    @app.delete("/api/line-groups/<lid>")
    @login_required
    def api_delete_line_group(lid: str):
        get_db().execute("DELETE FROM line_groups WHERE id=?", (lid,))
        get_db().commit()
        return jsonify({"ok": True})

    @app.post("/api/line/webhook")
    def line_webhook():
        body = request.get_json(silent=True) or {}
        events = body.get("events") or []
        conn = get_db()
        for ev in events:
            source = ev.get("source") or {}
            gid = source.get("groupId")
            if not gid:
                continue
            if conn.execute("SELECT 1 FROM line_groups WHERE group_id=?", (gid,)).fetchone():
                continue
            conn.execute(
                "INSERT INTO line_groups(id, name, group_id, message_type, created_at) VALUES (?,?,?,?,?)",
                (str(uuid.uuid4()), f"LINE {gid[-6:]}", gid, "summary", utc_now()),
            )
        conn.commit()
        return jsonify({"ok": True})

    # ---------- reports ----------
    @app.get("/api/reports/by-group")
    @login_required
    def report_by_group():
        mode = request.args.get("mode") or "day"
        items = query_donations_from_args(mode)
        groups = load_groups()
        member_to_group = {}
        for g in groups:
            for code in g["members"]:
                member_to_group[code] = g["id"]
        buckets: dict[str, dict[str, Any]] = {
            str(g["id"]): {
                "groupName": g["name"],
                "records": [],
                "totalPieces": 0,
                "totalWeight": 0.0,
                "totalRecords": 0,
            }
            for g in groups
        }
        ungrouped = {
            "groupName": "ไม่ได้จัดกลุ่ม",
            "records": [],
            "totalPieces": 0,
            "totalWeight": 0.0,
            "totalRecords": 0,
        }
        for rec in items:
            key = member_to_group.get(rec["branchCode"])
            bucket = buckets.get(str(key), ungrouped) if key else ungrouped
            bucket["records"].append(rec)
            bucket["totalPieces"] += rec["pieces"]
            bucket["totalWeight"] += rec["weightKg"]
            bucket["totalRecords"] += 1
        result = list(buckets.values())
        if ungrouped["totalRecords"]:
            result.append(ungrouped)
        return jsonify({"groups": result, "totals": totals_of(items)})

    @app.get("/api/reports/by-branch")
    @login_required
    def report_by_branch():
        year = int(request.args.get("year") or datetime.now().year)
        month = int(request.args.get("month") or datetime.now().month)
        items = donations_in_month(year, month)
        by: dict[str, dict[str, Any]] = {}
        for rec in items:
            b = by.setdefault(
                rec["branchCode"],
                {
                    "branchCode": rec["branchCode"],
                    "storeName": rec["storeName"],
                    "count": 0,
                    "pieces": 0,
                    "weight": 0.0,
                    "baskets": 0,
                },
            )
            b["count"] += 1
            b["pieces"] += rec["pieces"]
            b["weight"] += rec["weightKg"]
            b["baskets"] += rec["baskets"] or 0
        rows = sorted(by.values(), key=lambda x: x["branchCode"])
        return jsonify(
            {
                "year": year,
                "month": month,
                "label": month_label(year, month),
                "rows": rows,
                "totals": {
                    **totals_of(items),
                    "branchCount": len(rows),
                    "totalBaskets": sum(r["baskets"] for r in rows),
                },
            }
        )

    def query_donations_from_args(mode: str) -> list[dict[str, Any]]:
        if mode == "month":
            year = int(request.args.get("year") or datetime.now().year)
            month = int(request.args.get("month") or datetime.now().month)
            return donations_in_month(year, month)
        if mode == "range":
            start = request.args.get("start") or iso_today()
            end = request.args.get("end") or start
            return donations_in_range(start, end)
        day = request.args.get("date") or iso_today()
        return donations_in_range(day, day)

    @app.get("/api/export/donations.xlsx")
    @login_required
    def export_donations():
        mode = request.args.get("mode") or "day"
        items = query_donations_from_args(mode)
        wb = Workbook()
        ws = wb.active
        ws.title = "รายการบริจาค"
        ws.append(["วันที่", "รหัสสาขา", "ชื่อร้าน", "จำนวนชิ้น", "น้ำหนัก (กก.)", "แหล่ง"])
        for rec in items:
            ws.append(
                [
                    rec["pickupDate"] or rec["date"],
                    rec["branchCode"],
                    rec["storeName"],
                    rec["pieces"],
                    rec["weightKg"],
                    rec["source"],
                ]
            )
        t = totals_of(items)
        ws.append(["รวม", "", f"{t['totalRecords']} รายการ", t["totalPieces"], t["totalWeight"], ""])
        return xlsx_response(wb, "donations.xlsx")

    @app.get("/api/export/groups.xlsx")
    @login_required
    def export_groups():
        mode = request.args.get("mode") or "day"
        items = query_donations_from_args(mode)
        groups = load_groups()
        member_to_group = {}
        for g in groups:
            for code in g["members"]:
                member_to_group[code] = g["id"]
        by_id: dict[int, list] = {g["id"]: [] for g in groups}
        ungrouped: list = []
        for rec in items:
            gid = member_to_group.get(rec["branchCode"])
            if gid:
                by_id[gid].append(rec)
            else:
                ungrouped.append(rec)
        wb = Workbook()
        summary = wb.active
        summary.title = "สรุปตามกลุ่ม"
        summary.append(["กลุ่ม", "จำนวนครั้ง", "จำนวนชิ้น", "น้ำหนัก (กก.)"])
        for g in groups:
            recs = by_id[g["id"]]
            summary.append(
                [
                    g["name"],
                    len(recs),
                    sum(r["pieces"] for r in recs),
                    round(sum(r["weightKg"] for r in recs), 2),
                ]
            )
        if ungrouped:
            summary.append(
                [
                    "ไม่ได้จัดกลุ่ม",
                    len(ungrouped),
                    sum(r["pieces"] for r in ungrouped),
                    round(sum(r["weightKg"] for r in ungrouped), 2),
                ]
            )
        t = totals_of(items)
        summary.append(["รวมทั้งหมด", t["totalRecords"], t["totalPieces"], t["totalWeight"]])

        def add_sheet(title: str, recs: list[dict[str, Any]]) -> None:
            ws = wb.create_sheet(title[:31] or "กลุ่ม")
            ws.append(["ลำดับ", "วันที่", "รหัสสาขา", "ชื่อร้าน", "จำนวนชิ้น", "น้ำหนัก (กก.)"])
            for i, rec in enumerate(recs, 1):
                ws.append(
                    [
                        i,
                        rec["pickupDate"] or rec["date"],
                        rec["branchCode"],
                        rec["storeName"],
                        rec["pieces"],
                        rec["weightKg"],
                    ]
                )
            ws.append(
                [
                    "",
                    "",
                    "",
                    "รวม",
                    sum(r["pieces"] for r in recs),
                    round(sum(r["weightKg"] for r in recs), 2),
                ]
            )

        for g in groups:
            recs = by_id[g["id"]]
            if recs:
                add_sheet(g["name"], recs)
        if ungrouped:
            add_sheet("ไม่ได้จัดกลุ่ม", ungrouped)
        return xlsx_response(wb, "group-report.xlsx")

    @app.get("/api/export/branches.xlsx")
    @login_required
    def export_branches():
        year = int(request.args.get("year") or datetime.now().year)
        month = int(request.args.get("month") or datetime.now().month)
        items = donations_in_month(year, month)
        by: dict[str, dict[str, Any]] = {}
        for rec in items:
            b = by.setdefault(
                rec["branchCode"],
                {
                    "branchCode": rec["branchCode"],
                    "storeName": rec["storeName"],
                    "count": 0,
                    "pieces": 0,
                    "weight": 0.0,
                    "baskets": 0,
                },
            )
            b["count"] += 1
            b["pieces"] += rec["pieces"]
            b["weight"] += rec["weightKg"]
            b["baskets"] += rec["baskets"] or 0
        wb = Workbook()
        ws = wb.active
        ws.title = "สรุปตามสาขา"
        ws.append(["รหัสสาขา", "ชื่อร้าน", "จำนวนครั้ง", "จำนวนชิ้นรวม", "น้ำหนัก (กก.)", "จำนวนตะกร้า"])
        for r in sorted(by.values(), key=lambda x: x["branchCode"]):
            ws.append(
                [r["branchCode"], r["storeName"], r["count"], r["pieces"], round(r["weight"], 2), r["baskets"]]
            )
        t = totals_of(items)
        ws.append(
            [
                "รวมทั้งหมด",
                "",
                t["totalRecords"],
                t["totalPieces"],
                t["totalWeight"],
                sum(r["baskets"] for r in by.values()),
            ]
        )
        return xlsx_response(wb, f"branch-summary-{year}-{month:02d}.xlsx")

    def totals_of(items: list[dict[str, Any]]) -> dict[str, Any]:
        return {
            "totalPieces": sum(r["pieces"] for r in items),
            "totalWeight": round(sum(r["weightKg"] for r in items), 2),
            "totalRecords": len(items),
        }

    return app


if __name__ == "__main__":
    app = create_app()
    host = os.environ.get("GIFT_BIND", "0.0.0.0")
    port = int(os.environ.get("GIFT_PORT", "8110"))
    app.run(host=host, port=port, debug=os.environ.get("FLASK_DEBUG") == "1")
